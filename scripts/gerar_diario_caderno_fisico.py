from datetime import datetime
from pathlib import Path
import subprocess

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = Path('/Users/camila/Documents/Aquiles com Orion/Diario de Bordo Orion Academy para Caderno Fisico.docx')
REPOSITORY = 'https://github.com/Camis-Car/Orion-Academy'
GIT = '/Users/camila/.cache/codex-runtimes/codex-primary-runtime/dependencies/bin/fallback/git'


def set_cell_shading(cell, fill):
    properties = cell._tc.get_or_add_tcPr()
    shading = OxmlElement('w:shd')
    shading.set(qn('w:fill'), fill)
    properties.append(shading)


def set_cell_margins(cell, top=110, start=110, bottom=110, end=110):
    properties = cell._tc.get_or_add_tcPr()
    margins = properties.first_child_found_in('w:tcMar')
    if margins is None:
        margins = OxmlElement('w:tcMar')
        properties.append(margins)
    for side, value in {'top': top, 'start': start, 'bottom': bottom, 'end': end}.items():
        node = margins.find(qn(f'w:{side}'))
        if node is None:
            node = OxmlElement(f'w:{side}')
            margins.append(node)
        node.set(qn('w:w'), str(value))
        node.set(qn('w:type'), 'dxa')


def set_repeat_table_header(row):
    properties = row._tr.get_or_add_trPr()
    header = OxmlElement('w:tblHeader')
    header.set(qn('w:val'), 'true')
    properties.append(header)


def set_font(run, name='Aptos', size=11, bold=False, color=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn('w:ascii'), name)
    run._element.rPr.rFonts.set(qn('w:hAnsi'), name)
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)


def add_text(paragraph, text, **kwargs):
    run = paragraph.add_run(text)
    set_font(run, **kwargs)
    return run


def add_heading(document, text, level=1):
    paragraph = document.add_paragraph(style=f'Heading {level}')
    paragraph.paragraph_format.space_before = Pt(14 if level == 1 else 9)
    paragraph.paragraph_format.space_after = Pt(5)
    run = paragraph.add_run(text)
    set_font(run, name='Aptos Display', size=16 if level == 1 else 13, bold=True, color=(0, 0, 0))
    return paragraph


def get_log():
    result = subprocess.run(
        [GIT, 'log', '--all', '--date=iso-strict', '--pretty=format:%H%x09%ad%x09%an%x09%s', '--reverse'],
        cwd=ROOT,
        check=True,
        text=True,
        capture_output=True,
    )
    entries = []
    for line in result.stdout.splitlines():
        commit, timestamp, author, subject = line.split('\t', 3)
        moment = datetime.fromisoformat(timestamp)
        entries.append({'commit': commit, 'moment': moment, 'author': author, 'subject': subject})
    return entries


def write_document(entries):
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

    normal = document.styles['Normal']
    normal.font.name = 'Aptos'
    normal._element.rPr.rFonts.set(qn('w:ascii'), 'Aptos')
    normal._element.rPr.rFonts.set(qn('w:hAnsi'), 'Aptos')
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(5)

    for name in ('Title', 'Subtitle', 'Heading 1', 'Heading 2'):
        style = document.styles[name]
        style.font.color.rgb = RGBColor(0, 0, 0)

    title = document.add_paragraph(style='Title')
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(6)
    add_text(title, 'Diário de Bordo Orion Academy', name='Aptos Display', size=24, bold=True, color=(0, 0, 0))

    subtitle = document.add_paragraph(style='Subtitle')
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(20)
    add_text(subtitle, 'Guia cronológico para registro no caderno físico', name='Aptos', size=12, color=(70, 70, 70))

    intro = document.add_paragraph()
    intro.paragraph_format.space_after = Pt(10)
    add_text(intro, 'Finalidade. ', bold=True)
    add_text(intro, 'Este documento organiza as atividades registradas no histórico do projeto entre 2 de agosto e 6 de setembro de 2026. Use-o como referência para o diário físico, mantendo os registros futuros feitos no próprio dia.')

    note = document.add_paragraph()
    note.paragraph_format.space_after = Pt(13)
    add_text(note, 'Integridade do registro. ', bold=True)
    add_text(note, 'A hora indicada é a do commit no Git, isto é, o momento verificável de registro da alteração. Não foram inventados local, duração da sessão, observações, testes ou resultados; esses campos devem ser preenchidos à mão somente quando houver memória ou evidência confiável.')

    metadata = document.add_table(rows=4, cols=2)
    metadata.style = 'Table Grid'
    fields = [('Projeto', 'Orion Academy'), ('Período', '02/08/2026 a 06/09/2026'), ('Fuso horário', 'UTC−03:00'), ('Fonte das datas e horas', f'Histórico Git do repositório {REPOSITORY}')]
    for row, (label, value) in zip(metadata.rows, fields):
        for cell in row.cells:
            set_cell_margins(cell)
        set_cell_shading(row.cells[0], 'D9E3F0')
        paragraph = row.cells[0].paragraphs[0]
        add_text(paragraph, label, bold=True)
        add_text(row.cells[1].paragraphs[0], value)
    document.add_paragraph()

    add_heading(document, 'Como registrar no caderno físico', 1)
    instructions = [
        'Abra cada dia com a data, o local e os participantes. Se o local não estiver comprovado, escreva “local não registrado no histórico digital”.',
        'Copie a hora e a atividade exatamente como aparecem abaixo. Depois acrescente, em suas palavras, a decisão, a dúvida, o erro, o teste ou o resultado que você realmente se lembrar e puder confirmar.',
        'Cole ou anote o código curto do commit como evidência. No GitHub, cada código permite consultar o conteúdo alterado.',
        'Para sessões futuras, registre no momento em que elas ocorrerem: objetivo, procedimento, resultado, limitações e próximo passo.',
    ]
    for instruction in instructions:
        paragraph = document.add_paragraph(style='List Bullet')
        add_text(paragraph, instruction)

    add_heading(document, 'Registros cronológicos', 1)
    current_date = None
    weekdays = ['segunda-feira', 'terça-feira', 'quarta-feira', 'quinta-feira', 'sexta-feira', 'sábado', 'domingo']
    for entry in entries:
        moment = entry['moment']
        if moment.date() != current_date:
            current_date = moment.date()
            day_name = weekdays[moment.weekday()].capitalize()
            add_heading(document, f'{day_name}, {moment.strftime("%d/%m/%Y")}', 1)
            place = document.add_paragraph()
            place.paragraph_format.space_after = Pt(8)
            add_text(place, 'Local: ', bold=True)
            add_text(place, '__________________________________    ')
            add_text(place, 'Participantes: ', bold=True)
            add_text(place, '__________________________________')

        paragraph = document.add_paragraph()
        paragraph.paragraph_format.left_indent = Inches(0.14)
        paragraph.paragraph_format.first_line_indent = Inches(-0.14)
        paragraph.paragraph_format.space_after = Pt(2)
        add_text(paragraph, f'{moment.strftime("%H:%M:%S")} — ', bold=True)
        add_text(paragraph, entry['subject'])
        evidence = document.add_paragraph()
        evidence.paragraph_format.left_indent = Inches(0.18)
        evidence.paragraph_format.space_after = Pt(7)
        add_text(evidence, f'Evidência: commit {entry["commit"][:7]} | autoria registrada: {entry["author"]}', size=9, color=(92, 92, 92))

    add_heading(document, 'Registros para as próximas sessões', 1)
    paragraph = document.add_paragraph()
    add_text(paragraph, 'Preencha esta estrutura no caderno no dia em que a atividade acontecer.')
    table = document.add_table(rows=1, cols=5)
    table.style = 'Table Grid'
    headers = ['Data e hora', 'Local e participantes', 'Objetivo e procedimento', 'Resultado ou dúvida', 'Próximo passo e evidência']
    for cell, header in zip(table.rows[0].cells, headers):
        set_cell_shading(cell, '173A5E')
        set_cell_margins(cell, top=130, start=100, bottom=130, end=100)
        paragraph = cell.paragraphs[0]
        add_text(paragraph, header, size=9, bold=True, color=(255, 255, 255))
    set_repeat_table_header(table.rows[0])
    for _ in range(3):
        row = table.add_row()
        for cell in row.cells:
            set_cell_margins(cell, top=400, start=100, bottom=400, end=100)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_text(footer, 'Diário de Bordo Orion Academy | Guia para caderno físico', size=8, color=(100, 100, 100))
    document.save(OUTPUT)


if __name__ == '__main__':
    write_document(get_log())
    print(OUTPUT)
